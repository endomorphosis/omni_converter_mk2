"""
Simple test for XLSX processor functionality.
"""

import io
import sys
from deprecated.processors.xlsx_processor import XlsxProcessor, OPENPYXL_AVAILABLE

def main():
    """Run a simple test of the XLSX processor."""
    print("Testing XLSX processor...")
    
    # Check if openpyxl is available
    if not OPENPYXL_AVAILABLE:
        print("Warning: openpyxl is not available. XLSX processing will be limited.")
        return
    
    # Create a processor
    processor = XlsxProcessor()
    print(f"Created processor: {processor.__class__.__name__}")
    print(f"Supported formats: {processor.supported_formats}")
    print(f"Processor info: {processor.get_processor_info()}")
    
    # Create a simple test XLSX
    try:
        import openpyxl
        
        # Create a new workbook
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Test Sheet"
        
        # Add some data
        ws['A1'] = "ID"
        ws['B1'] = "Name"
        ws['C1'] = "Value"
        
        ws['A2'] = 1
        ws['B2'] = "Item 1"
        ws['C2'] = 100.5
        
        ws['A3'] = 2
        ws['B3'] = "Item 2"
        ws['C3'] = 200.75
        
        # Save to BytesIO
        xlsx_data = io.BytesIO()
        wb.save(xlsx_data)
        xlsx_bytes = xlsx_data.getvalue()
        
        print("\nProcessing test XLSX...")
        options = {
            'include_empty_cells': True,
            'max_rows': 100
        }
        
        # Test text extraction
        text = processor.extract_text(xlsx_bytes, options)
        print("\nExtracted text preview:")
        print(text[:200] + "..." if len(text) > 200 else text)
        
        # Test metadata extraction
        metadata = processor.extract_metadata(xlsx_bytes, options)
        print("\nExtracted metadata preview:")
        print(str(metadata)[:200] + "..." if len(str(metadata)) > 200 else str(metadata))
        
        # Test structure extraction
        structure = processor.extract_structure(xlsx_bytes, options)
        print("\nExtracted structure preview:")
        print(str(structure)[:200] + "..." if len(str(structure)) > 200 else str(structure))
        
        # Test full processing
        text, metadata, sections = processor.process_document(xlsx_bytes, options)
        print("\nFull processing successful!")
        
        print("\nXLSX processor test completed successfully!")
        
    except Exception as e:
        print(f"Error during testing: {e}")
        raise

if __name__ == "__main__":
    main()