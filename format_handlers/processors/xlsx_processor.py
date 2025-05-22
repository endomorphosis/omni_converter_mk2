"""
XLSX processor implementation using openpyxl.

This module provides a concrete implementation of DocumentProcessor for XLSX files
using the openpyxl library.
"""

import io
from typing import Any, Dict, List, Tuple, Optional, BinaryIO
from datetime import datetime

from format_handlers.processors.document_processor import DocumentProcessor
from utils.logger import logger

try:
    import openpyxl
    from openpyxl.workbook import Workbook
    OPENPYXL_AVAILABLE = True
except ImportError:
    logger.warning("openpyxl not available, XLSX processing will not be available")
    OPENPYXL_AVAILABLE = False


class XlsxProcessor(DocumentProcessor):
    """
    XLSX processor implementation using openpyxl.
    
    This class provides functionality to extract text, metadata, and structure
    from XLSX files using the openpyxl library.
    """
    
    def __init__(self):
        """Initialize the XLSX processor."""
        self._supported_formats = ["xlsx"]
    
    def can_process(self, format_name: str) -> bool:
        """
        Check if this processor can handle the given format.
        
        Args:
            format_name: The name of the format to check.
            
        Returns:
            True if this processor can handle the format and openpyxl is available,
            False otherwise.
        """
        return OPENPYXL_AVAILABLE and format_name.lower() in self._supported_formats
    
    @property
    def supported_formats(self) -> List[str]:
        """
        Get the list of formats supported by this processor.
        
        Returns:
            A list of format names supported by this processor.
        """
        return self._supported_formats if OPENPYXL_AVAILABLE else []
    
    def get_processor_info(self) -> Dict[str, Any]:
        """
        Get information about this processor.
        
        Returns:
            A dictionary containing information about this processor.
        """
        info = {
            "name": "XlsxProcessor",
            "supported_formats": self._supported_formats,
            "available": OPENPYXL_AVAILABLE
        }
        
        if OPENPYXL_AVAILABLE:
            info["version"] = getattr(openpyxl, "__version__", "Unknown")
        
        return info
    
    def extract_text(self, data: bytes, options: Dict[str, Any]) -> str:
        """
        Extract plain text from an XLSX document.
        
        Args:
            data: The binary data of the XLSX document.
            options: Processing options.
                include_empty_cells: Whether to include empty cells (default: False)
                max_rows: Maximum number of rows to extract per sheet (default: 1000)
                
        Returns:
            Extracted text from the XLSX document.
            
        Raises:
            ValueError: If openpyxl is not available or the data cannot be processed as an XLSX.
        """
        if not OPENPYXL_AVAILABLE:
            raise ValueError("openpyxl is not available for XLSX text extraction")
        
        try:
            # Create a file-like object from the bytes
            xlsx_file = io.BytesIO(data)
            
            # Open the XLSX file
            wb = openpyxl.load_workbook(xlsx_file, read_only=True, data_only=True)
            
            # Get options
            include_empty_cells = options.get('include_empty_cells', False)
            max_rows = options.get('max_rows', 1000)
            
            # Extract text from each sheet
            sheet_texts = []
            
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                
                # Add sheet header
                sheet_texts.append(f"\n\n--- Sheet: {sheet_name} ---\n")
                
                # Extract data from cells
                rows = []
                row_count = 0
                
                for row in ws.iter_rows(max_row=max_rows):
                    row_count += 1
                    cells = []
                    
                    for cell in row:
                        value = cell.value
                        
                        # Format the value as a string
                        if value is None:
                            if include_empty_cells:
                                cells.append("")
                            else:
                                # Skip empty cells if not including them
                                continue
                        elif isinstance(value, datetime):
                            cells.append(value.strftime("%Y-%m-%d %H:%M:%S"))
                        else:
                            cells.append(str(value))
                    
                    # Skip empty rows
                    if cells:
                        rows.append("\t".join(cells))
                
                # Add rows to sheet text
                sheet_texts.append("\n".join(rows))
                
                # Add notice if we hit the row limit
                if row_count >= max_rows:
                    sheet_texts.append(f"\n[Row limit of {max_rows} reached. Additional rows not shown.]")
            
            # Join all sheets
            return "\n".join(sheet_texts)
            
        except Exception as e:
            logger.error(f"Error extracting text from XLSX: {e}")
            raise ValueError(f"Error extracting text from XLSX: {e}")
    
    def extract_metadata(self, data: bytes, options: Dict[str, Any]) -> Dict[str, Any]:
        """
        Extract metadata from an XLSX document.
        
        Args:
            data: The binary data of the XLSX document.
            options: Processing options.
            
        Returns:
            Metadata extracted from the XLSX document.
            
        Raises:
            ValueError: If openpyxl is not available or the data cannot be processed as an XLSX.
        """
        if not OPENPYXL_AVAILABLE:
            raise ValueError("openpyxl is not available for XLSX metadata extraction")
        
        try:
            # Create a file-like object from the bytes
            xlsx_file = io.BytesIO(data)
            
            # Open the XLSX file
            wb = openpyxl.load_workbook(xlsx_file, read_only=True)
            
            # Extract document info
            metadata = {
                "file_size_bytes": len(data),
                "sheet_count": len(wb.sheetnames),
                "sheets": wb.sheetnames
            }
            
            # Extract document properties
            if hasattr(wb, "properties"):
                props = wb.properties
                if props.title:
                    metadata["title"] = props.title
                if props.creator:
                    metadata["creator"] = props.creator
                if props.subject:
                    metadata["subject"] = props.subject
                if props.keywords:
                    metadata["keywords"] = props.keywords
                if props.category:
                    metadata["category"] = props.category
                if props.description:
                    metadata["description"] = props.description
                if props.created:
                    metadata["creation_date"] = props.created.isoformat() if hasattr(props.created, "isoformat") else str(props.created)
                if props.modified:
                    metadata["modification_date"] = props.modified.isoformat() if hasattr(props.modified, "isoformat") else str(props.modified)
                if props.lastModifiedBy:
                    metadata["last_modified_by"] = props.lastModifiedBy
                if props.revision:
                    metadata["revision"] = props.revision
            
            # Extract sheet statistics
            sheet_stats = []
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                
                # Get dimensions if available
                dimensions = "Unknown"
                if hasattr(ws, "calculate_dimension") and callable(getattr(ws, "calculate_dimension")):
                    try:
                        dimensions = ws.calculate_dimension()
                    except:
                        # Some sheets may not have data
                        dimensions = "Empty or Error"
                
                sheet_stats.append({
                    "name": sheet_name,
                    "dimensions": dimensions,
                    "sheet_state": ws.sheet_state
                })
            
            metadata["sheet_statistics"] = sheet_stats
            
            return metadata
            
        except Exception as e:
            logger.error(f"Error extracting metadata from XLSX: {e}")
            raise ValueError(f"Error extracting metadata from XLSX: {e}")
    
    def extract_structure(self, data: bytes, options: Dict[str, Any]) -> List[Dict[str, Any]]:
        """
        Extract structural elements from an XLSX document.
        
        Args:
            data: The binary data of the XLSX document.
            options: Processing options.
            
        Returns:
            A list of structural elements extracted from the XLSX document.
            
        Raises:
            ValueError: If openpyxl is not available or the data cannot be processed as an XLSX.
        """
        if not OPENPYXL_AVAILABLE:
            raise ValueError("openpyxl is not available for XLSX structure extraction")
        
        try:
            # Create a file-like object from the bytes
            xlsx_file = io.BytesIO(data)
            
            # Open the XLSX file
            wb = openpyxl.load_workbook(xlsx_file, read_only=True, data_only=True)
            
            # Extract structure
            structure = []
            
            # Add workbook as a section
            structure.append({
                "type": "workbook",
                "content": "XLSX Workbook"
            })
            
            # Extract sheets and sample data
            max_rows = options.get('structure_max_rows', 20)
            max_columns = options.get('structure_max_cols', 10)
            
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                
                # Sample data for display
                sample_data = []
                
                row_count = 0
                for row in ws.iter_rows(max_row=max_rows, max_col=max_columns):
                    row_data = []
                    for cell in row:
                        value = cell.value
                        if isinstance(value, datetime):
                            value = value.strftime("%Y-%m-%d %H:%M:%S")
                        row_data.append(str(value) if value is not None else "")
                    
                    sample_data.append(row_data)
                    row_count += 1
                
                # Add sheet structure
                sheet_structure = {
                    "type": "sheet",
                    "name": sheet_name,
                    "content": {
                        "sample_data": sample_data
                    }
                }
                
                # Add dimensions if available
                if hasattr(ws, "calculate_dimension") and callable(getattr(ws, "calculate_dimension")):
                    try:
                        sheet_structure["content"]["dimensions"] = ws.calculate_dimension()
                    except:
                        sheet_structure["content"]["dimensions"] = "Empty or Error"
                
                # Add sample size information
                if row_count >= max_rows:
                    sheet_structure["content"]["note"] = f"Showing {max_rows} rows, {max_columns} columns (sample)"
                
                structure.append(sheet_structure)
                
                # Add named ranges if any are defined
                if hasattr(wb, "defined_names") and wb.defined_names:
                    defined_names = []
                    for name in wb.defined_names:
                        if name.name.startswith('_'):  # Internal name
                            continue
                        defined_names.append({
                            "name": name.name,
                            "destinations": list(name.destinations)
                        })
                    
                    if defined_names:
                        structure.append({
                            "type": "named_ranges",
                            "content": defined_names
                        })
            
            return structure
            
        except Exception as e:
            logger.error(f"Error extracting structure from XLSX: {e}")
            raise ValueError(f"Error extracting structure from XLSX: {e}")
    
    def process_document(self, data: bytes, options: Dict[str, Any]) -> Tuple[str, Dict[str, Any], List[Dict[str, Any]]]:
        """
        Process an XLSX document completely, extracting text, metadata, and structure.
        
        Args:
            data: The binary data of the XLSX document.
            options: Processing options.
            
        Returns:
            A tuple of (text content, metadata, sections).
            
        Raises:
            ValueError: If openpyxl is not available or the data cannot be processed as an XLSX.
        """
        if not OPENPYXL_AVAILABLE:
            raise ValueError("openpyxl is not available for XLSX processing")
        
        try:
            # Extract text, metadata, and structure
            text = self.extract_text(data, options)
            metadata = self.extract_metadata(data, options)
            sections = self.extract_structure(data, options)
            
            # Create a human-readable text version
            text_content = [f"XLSX Document: {metadata.get('title', 'Untitled')}"]
            
            if "creator" in metadata:
                text_content.append(f"Creator: {metadata['creator']}")
            
            if "subject" in metadata:
                text_content.append(f"Subject: {metadata['subject']}")
            
            if "creation_date" in metadata:
                text_content.append(f"Created: {metadata['creation_date']}")
            
            if "sheet_count" in metadata:
                text_content.append(f"Total Sheets: {metadata['sheet_count']}")
                text_content.append(f"Sheet Names: {', '.join(metadata['sheets'])}")
            
            text_content.append("\n--- Spreadsheet Content ---\n")
            text_content.append(text)
            
            return "\n".join(text_content), metadata, sections
            
        except Exception as e:
            logger.error(f"Error processing XLSX document: {e}")
            raise ValueError(f"Error processing XLSX document: {e}")


# Create a global instance for usage
xlsx_processor = XlsxProcessor()
