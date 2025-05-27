"""
DOCX processor implementation using python-docx.

This module provides a concrete implementation of DocumentProcessor for DOCX files
using the python-docx library.
"""

import io
from typing import Any, Dict, List, Tuple, Optional, BinaryIO
from datetime import datetime

from format_handlers.processors.document_processor import DocumentProcessor
from logger import logger
from utils.common.try_except_decorator import try_except

try:
    import docx
    from docx.document import Document as DocxDocument
    PYTHON_DOCX_AVAILABLE = True
except ImportError:
    logger.warning("python-docx not available, DOCX processing will not be available")
    PYTHON_DOCX_AVAILABLE = False


def set_metadata(props: object, metadata: Dict[str, Any]) -> None:
    list_of_attr = [attr for attr in dir(props)]
    for key in metadata.keys():
        if key in list_of_attr:
            value = getattr(props, key)
            match value:
                case datetime():
                    metadata[key] = value.isoformat()
                case str():
                    metadata[key] = value
                case _: # Stringify all other types
                    metadata[key] = str(value)


class DocxProcessor(DocumentProcessor):
    """
    DOCX processor implementation using python-docx.
    
    This class provides functionality to extract text, metadata, and structure
    from DOCX files using the python-docx library.
    """
    
    def __init__(self):
        """Initialize the DOCX processor."""
        self._supported_formats = ["docx"]
    
    def can_process(self, format_name: str) -> bool:
        """
        Check if this processor can handle the given format.
        
        Args:
            format_name: The name of the format to check.
            
        Returns:
            True if this processor can handle the format and python-docx is available,
            False otherwise.
        """
        return PYTHON_DOCX_AVAILABLE and format_name.lower() in self.supported_formats
    
    @property
    def supported_formats(self) -> List[str]:
        """
        Get the list of formats supported by this processor.
        
        Returns:
            A list of format names supported by this processor.
        """
        return self._supported_formats if PYTHON_DOCX_AVAILABLE else []
    
    def get_processor_info(self) -> Dict[str, Any]:
        """
        Get information about this processor.
        
        Returns:
            A dictionary containing information about this processor.
        """
        info = {
            "name": "DocxProcessor",
            "supported_formats": self.supported_formats,
            "available": PYTHON_DOCX_AVAILABLE
        }
        
        if PYTHON_DOCX_AVAILABLE:
            info["version"] = getattr(docx, "__version__", "Unknown")
        
        return info
    
    @try_except(raise_=True, raise_as=ValueError, msg="Error extracting text from DOCX")
    def extract_text(self, data: bytes, options: Dict[str, Any]) -> str:
        """
        Extract plain text from a DOCX document.
        
        Args:
            data: The binary data of the DOCX document.
            options: Processing options.
                
        Returns:
            Extracted text from the DOCX document.
            
        Raises:
            ValueError: If python-docx is not available or the data cannot be processed as a DOCX.
        """
        if not PYTHON_DOCX_AVAILABLE:
            raise ValueError("python-docx is not available for DOCX text extraction")

        # Create a file-like object from the bytes
        docx_file = io.BytesIO(data)
        
        # Open the DOCX file
        doc = docx.Document(docx_file)
        
        # Extract text from each paragraph
        paragraphs = []
        
        for para in doc.paragraphs:
            # Skip empty paragraphs
            if para.text.strip():
                paragraphs.append(para.text)
        
        # Extract text from tables if requested
        include_tables = options.get('include_tables', True)
        if include_tables:
            for table in doc.tables:
                for row in table.rows:
                    cells = []
                    for cell in row.cells:
                        # Get text from each cell
                        if cell.text.strip():
                            cells.append(cell.text.strip())
                    
                    if cells:
                        paragraphs.append(" | ".join(cells))
        
        # Join all paragraphs
        return "\n\n".join(paragraphs)


    @try_except(raise_=True, raise_as=ValueError, msg="Error extracting metadata from DOCX")
    def extract_metadata(self, data: bytes, options: Dict[str, Any]) -> Dict[str, Any]:
        """
        Extract metadata from a DOCX document.
        
        Args:
            data: The binary data of the DOCX document.
            options: Processing options.
            
        Returns:
            Metadata extracted from the DOCX document.
            
        Raises:
            ValueError: If python-docx is not available or the data cannot be processed as a DOCX.
        """
        if not PYTHON_DOCX_AVAILABLE:
            raise ValueError("python-docx is not available for DOCX metadata extraction")

        # Create a file-like object from the bytes
        docx_file = io.BytesIO(data)
        
        # Open the DOCX file
        doc = docx.Document(docx_file)
        
        # Extract document info
        metadata = {
            "title": None,
            "author": None,
            "subject": None,
            "keywords": None,
            "category": None,
            "comments": None,
            "creation_date": None,
            "modification_date": None,
            "last_modified_by": None,
            "revision": None,
            "language": None,
            "file_size_bytes": len(data),
            "paragraph_count": len(doc.paragraphs),
            "table_count": len(doc.tables),
            "section_count": len(doc.sections)
        }

        # Extract core properties if available
        if hasattr(doc, "core_properties"):
            set_metadata(doc.core_properties, metadata)

        # Extract document statistics
        stats = {
            "character_count": 0,
            "word_count": 0,
            "paragraph_count": len(doc.paragraphs)
        }
        
        # Calculate character and word counts
        for para in doc.paragraphs:
            if para.text:
                stats["character_count"] += len(para.text)
                stats["word_count"] += len(para.text.split())
        
        metadata["statistics"] = stats

        return metadata

    @try_except(raise_=True, raise_as=ValueError, msg="Error extracting structure from DOCX")
    def extract_structure(self, data: bytes, options: Dict[str, Any]) -> List[Dict[str, Any]]:
        """
        Extract structural elements from a DOCX document.
        
        Args:
            data: The binary data of the DOCX document.
            options: Processing options.
            
        Returns:
            A list of structural elements extracted from the DOCX document.
            
        Raises:
            ValueError: If python-docx is not available or the data cannot be processed as a DOCX.
        """
        if not PYTHON_DOCX_AVAILABLE:
            raise ValueError("python-docx is not available for DOCX structure extraction")

        # Create a file-like object from the bytes
        docx_file = io.BytesIO(data)
        
        # Open the DOCX file
        doc = docx.Document(docx_file)
        
        # Extract structure
        structure = []
        
        # Add document as a section
        structure.append({
            "type": "document",
            "content": "DOCX Document"
        })
        
        # Extract headings and content
        current_heading = None
        heading_content = []
        
        for para in doc.paragraphs:
            # Check if this paragraph is a heading
            if para.style.name.startswith('Heading'):
                # If we had a previous heading, add it to the structure
                if current_heading and heading_content:
                    structure.append({
                        "type": "section",
                        "heading": current_heading,
                        "level": int(current_heading.split(':')[0].replace('Heading', '')) if current_heading[0].isdigit() else 1,
                        "content": "\n".join(heading_content)
                    })
                
                # Start a new heading
                current_heading = para.style.name + ": " + para.text
                heading_content = []
            elif para.text.strip():
                # Add content to the current heading, or to a default section if no heading yet
                if not current_heading:
                    current_heading = "Document"
                heading_content.append(para.text)
        
        # Add the last heading if there was one
        if current_heading and heading_content:
            structure.append({
                "type": "section",
                "heading": current_heading,
                "content": "\n".join(heading_content)
            })
        
        # Extract tables
        for i, table in enumerate(doc.tables):
            table_data = []
            for row in table.rows:
                row_data = []
                for cell in row.cells:
                    row_data.append(cell.text)
                table_data.append(row_data)
            
            structure.append({
                "type": "table",
                "table_number": i + 1,
                "rows": len(table.rows),
                "columns": len(table.rows[0].cells) if table.rows else 0,
                "content": table_data
            })
        
        # Extract images (limited info since we can't get image content easily)
        image_count = 0
        for rel in doc.part.rels.values():
            if "image" in rel.reltype:
                image_count += 1
        
        if image_count > 0:
            structure.append({
                "type": "images",
                "count": image_count,
                "content": f"Document contains {image_count} image(s)"
            })
        
        # Extract document properties
        if hasattr(doc, "sections"):
            for i, section in enumerate(doc.sections):
                section_info = {
                    "type": "document_section",
                    "section_number": i + 1,
                    "content": {}
                }
                
                if hasattr(section, "page_height") and hasattr(section, "page_width"):
                    section_info["content"]["page_size"] = {
                        "width": section.page_width.inches if hasattr(section.page_width, "inches") else "Unknown",
                        "height": section.page_height.inches if hasattr(section.page_height, "inches") else "Unknown"
                    }
                
                if hasattr(section, "orientation"):
                    section_info["content"]["orientation"] = section.orientation
                
                structure.append(section_info)
        
        return structure

    @try_except(raise_=True, raise_as=ValueError, msg="Error processing DOCX document")
    def process_document(self, data: bytes, options: Dict[str, Any]) -> Tuple[str, Dict[str, Any], List[Dict[str, Any]]]:
        """
        Process a DOCX document completely, extracting text, metadata, and structure.
        
        Args:
            data: The binary data of the DOCX document.
            options: Processing options.
            
        Returns:
            A tuple of (text content, metadata, sections).
            
        Raises:
            ValueError: If python-docx is not available or the data cannot be processed as a DOCX.
        """
        if not PYTHON_DOCX_AVAILABLE:
            raise ValueError("python-docx is not available for DOCX processing")

        # Extract text, metadata, and structure
        text = self.extract_text(data, options)
        metadata = self.extract_metadata(data, options)
        sections = self.extract_structure(data, options)
        
        # Create a human-readable text version
        text_content = [f"DOCX Document: {metadata.get('title', 'Untitled')}"]
        
        if "author" in metadata:
            text_content.append(f"Author: {metadata['author']}")
        
        if "subject" in metadata:
            text_content.append(f"Subject: {metadata['subject']}")
        
        if "creation_date" in metadata:
            text_content.append(f"Created: {metadata['creation_date']}")
        
        if "statistics" in metadata:
            stats = metadata["statistics"]
            text_content.append(f"Word Count: {stats.get('word_count', 'Unknown')}")
            text_content.append(f"Paragraphs: {stats.get('paragraph_count', 'Unknown')}")
        
        text_content.append("\n--- Document Text ---\n")
        text_content.append(text)
        
        return "\n".join(text_content), metadata, sections



# Create a global instance for usage
docx_processor = DocxProcessor()
